import html
import inspect
import io
import traceback

import openpyxl
import pandas as pd
import streamlit as st

from main import run_lqa

st.set_page_config(
    page_title="AI 翻译校对",
    page_icon="◇",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ---------------------------------------------------------------------------
# Theme — follow Streamlit light/dark tokens, do not hard-code white
# ---------------------------------------------------------------------------
def inject_theme():
    st.markdown(
        """
        <style>
        @import url("https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;550;600;700&display=swap");

        html, body, [class*="css"] {
            font-family: "Source Sans 3", "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif;
        }

        [data-testid="stHeader"] { background: transparent; }
        footer { visibility: hidden; }

        .block-container {
            padding-top: 1.25rem;
            padding-bottom: 3rem;
            max-width: 1200px;
        }

        /* Soft wash from theme primary so light mode is not flat white
           and dark mode keeps its native surface */
        .stApp {
            background:
                radial-gradient(
                    1200px 280px at 0% -10%,
                    color-mix(in srgb, var(--primary-color) 16%, var(--background-color)) 0%,
                    var(--background-color) 70%
                ) !important;
        }

        [data-testid="stSidebar"] {
            background: var(--secondary-background-color) !important;
            border-right: 1px solid color-mix(in srgb, var(--text-color) 12%, transparent);
        }
        [data-testid="stSidebar"] > div:first-child {
            padding-top: 1rem;
        }

        .side-title {
            font-size: 13px;
            font-weight: 650;
            color: var(--text-color);
            margin: 16px 0 8px 0;
            opacity: 0.88;
        }
        .side-title:first-child { margin-top: 0; }

        .hero {
            display: flex;
            align-items: flex-end;
            justify-content: space-between;
            gap: 16px;
            padding: 6px 0 16px 0;
            border-bottom: 1px solid color-mix(in srgb, var(--text-color) 12%, transparent);
            margin-bottom: 18px;
        }
        .hero-kicker {
            font-size: 11px;
            letter-spacing: 0.16em;
            text-transform: uppercase;
            color: var(--primary-color);
            font-weight: 650;
            margin-bottom: 6px;
        }
        .hero h1 {
            font-size: 26px;
            line-height: 1.15;
            margin: 0;
            color: var(--text-color);
            font-weight: 700;
        }
        .hero p {
            margin: 6px 0 0 0;
            color: var(--text-color);
            opacity: 0.62;
            font-size: 14px;
        }
        .hero-meta {
            color: var(--text-color);
            opacity: 0.45;
            font-size: 13px;
            text-align: right;
            white-space: nowrap;
        }

        .meta-line {
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            gap: 8px 0;
            color: var(--text-color);
            font-size: 13px;
            margin: 0 0 16px 0;
            opacity: 0.78;
        }
        .meta-line span { white-space: nowrap; }
        .meta-dot {
            width: 3px;
            height: 3px;
            border-radius: 50%;
            background: currentColor;
            opacity: 0.45;
            margin: 0 10px;
        }

        .empty-wrap {
            border: 1px dashed color-mix(in srgb, var(--text-color) 22%, transparent);
            background: color-mix(in srgb, var(--secondary-background-color) 80%, transparent);
            border-radius: 14px;
            padding: 56px 24px;
            text-align: center;
            color: var(--text-color);
            opacity: 0.85;
        }
        .empty-wrap h3 {
            color: var(--text-color);
            font-size: 18px;
            margin: 0 0 8px 0;
            opacity: 1;
        }

        /* Symmetric action buttons — slate blue, never theme-red */
        section[data-testid="stSidebar"] .stButton button,
        section[data-testid="stSidebar"] .stDownloadButton button {
            width: 100%;
            height: 42px;
            border-radius: 8px;
            font-weight: 600;
            font-size: 14px;
            box-shadow: none !important;
        }
        section[data-testid="stSidebar"] .stButton button[kind="primary"],
        section[data-testid="stSidebar"] .stDownloadButton button {
            background: #2b6cb0 !important;
            border: 1px solid #2b6cb0 !important;
            color: #ffffff !important;
        }
        section[data-testid="stSidebar"] .stButton button[kind="primary"]:hover,
        section[data-testid="stSidebar"] .stDownloadButton button:hover {
            background: #255f9c !important;
            border-color: #255f9c !important;
        }
        section[data-testid="stSidebar"] .stButton button[kind="primary"]:disabled,
        section[data-testid="stSidebar"] .stButton button:disabled {
            opacity: 0.45;
        }
        section[data-testid="stSidebar"] .stButton button[kind="secondary"] {
            background: transparent !important;
            color: var(--text-color) !important;
            border: 1px solid color-mix(in srgb, var(--text-color) 22%, transparent) !important;
        }
        section[data-testid="stSidebar"] .stButton button[kind="secondary"]:hover {
            background: color-mix(in srgb, var(--text-color) 6%, transparent) !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


inject_theme()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
SESSION_KEYS = (
    "wb",
    "orig_wb",
    "orig_name",
    "file_sig",
    "result_df",
    "result_sheet",
    "source_lang",
    "target_lang",
    "lqa_running",
    "lqa_error",
    "lqa_error_tb",
    "selected_sheet",
)


def clear_session_state():
    """Clears workbook and results from memory when file is removed or changed."""
    for key in SESSION_KEYS:
        st.session_state.pop(key, None)


def unique_headers(raw_headers):
    """Replace empty headers and de-duplicate colliding names."""
    seen = {}
    headers = []
    for i, h in enumerate(raw_headers):
        name = str(h).strip() if h is not None and str(h).strip() != "" else f"Unnamed_{i+1}"
        n = seen.get(name, 0)
        seen[name] = n + 1
        headers.append(name if n == 0 else f"{name}_{n+1}")
    return headers


def sheet_to_df(ws):
    data = list(ws.values)
    if not data:
        return pd.DataFrame()

    headers = unique_headers(data[0])
    df = pd.DataFrame(data[1:], columns=headers)
    return df.fillna("")


def option_index(options, value, default=0):
    try:
        return options.index(value)
    except (ValueError, IndexError):
        return default


def clone_workbook(wb):
    """Best-effort copy so a failed / repeated run cannot mutate the original."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def call_run_lqa(**kwargs):
    """Pass optional cancel callback only if run_lqa accepts it."""
    try:
        params = inspect.signature(run_lqa).parameters
    except (TypeError, ValueError):
        params = {}

    if "should_cancel" in params:
        kwargs["should_cancel"] = lambda: bool(st.session_state.get("lqa_cancel"))
    elif "cancel_check" in params:
        kwargs["cancel_check"] = lambda: bool(st.session_state.get("lqa_cancel"))

    return run_lqa(**kwargs)


def render_hero(subtitle="上传 Excel 文档，配置语言列后开始校对。"):
    st.markdown(
        f"""
        <div class="hero">
          <div>
            <div class="hero-kicker">Localization QA</div>
            <h1>AI 翻译校对</h1>
            <p>{html.escape(subtitle)}</p>
          </div>
          <div class="hero-meta">Game LQA</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_meta(parts):
    chunks = []
    for i, part in enumerate(parts):
        if i:
            chunks.append('<span class="meta-dot"></span>')
        chunks.append(f"<span>{html.escape(str(part))}</span>")
    st.markdown(f'<div class="meta-line">{"".join(chunks)}</div>', unsafe_allow_html=True)


def sidebar_heading(text):
    st.markdown(f'<div class="side-title">{html.escape(text)}</div>', unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Sidebar: document
# ---------------------------------------------------------------------------
with st.sidebar:
    sidebar_heading("文档")
    uploaded_file = st.file_uploader(
        "上传 Excel",
        type=["xlsx"],
        help="仅支持 .xlsx（openpyxl 无法读取旧版 .xls）。",
    )

if uploaded_file is None:
    clear_session_state()
    render_hero()
    st.markdown(
        """
        <div class="empty-wrap">
          <h3>还没有文档</h3>
          从左侧上传一份 .xlsx 本地化表即可开始。<br/>
          校对结果会写回原工作簿，其它工作表与格式会保留。
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.stop()

file_sig = (uploaded_file.name, uploaded_file.size, getattr(uploaded_file, "file_id", None))

if st.session_state.get("file_sig") != file_sig:
    clear_session_state()
    st.session_state.file_sig = file_sig
    st.session_state.orig_name = uploaded_file.name
    try:
        uploaded_file.seek(0)
        wb_loaded = openpyxl.load_workbook(uploaded_file, data_only=False)
        st.session_state.orig_wb = wb_loaded
        st.session_state.wb = clone_workbook(wb_loaded)
    except Exception as e:
        render_hero()
        st.error(f"文件读取错误: {e}")
        st.stop()

wb = st.session_state.wb
sheet_names = wb.sheetnames

# ---------------------------------------------------------------------------
# Sidebar: sheet + columns + glossary + actions
# ---------------------------------------------------------------------------
with st.sidebar:
    st.caption(uploaded_file.name)
    sidebar_heading("工作表")
    selected_sheet = st.selectbox(
        "工作表",
        sheet_names,
        index=option_index(sheet_names, st.session_state.get("selected_sheet")),
        label_visibility="collapsed",
    )

if st.session_state.get("selected_sheet") != selected_sheet:
    st.session_state.selected_sheet = selected_sheet
    if st.session_state.get("result_sheet") != selected_sheet:
        st.session_state.pop("result_df", None)

ws = wb[selected_sheet]
df = sheet_to_df(ws)

if df.empty:
    render_hero()
    st.warning("所选工作表为空，请在左侧更换工作表。")
    st.stop()

columns = df.columns.tolist()

with st.sidebar:
    sidebar_heading("语言列")
    source_lang = st.selectbox(
        "源语言列",
        columns,
        index=option_index(columns, st.session_state.get("source_lang"), 0),
    )
    target_lang = st.selectbox(
        "目标语言列",
        columns,
        index=option_index(
            columns,
            st.session_state.get("target_lang"),
            default=min(1, len(columns) - 1),
        ),
    )

    st.session_state.source_lang = source_lang
    st.session_state.target_lang = target_lang
    same_column = source_lang == target_lang
    if same_column:
        st.warning("源语言和目标语言不能是同一列。")

    sidebar_heading("校对设置")
    use_glossary = st.toggle("使用术语库", value=True)
    game_id = None
    if use_glossary:
        game_id = st.selectbox("项目 / 游戏 ID", ["S06"])

    sidebar_heading("操作")
    can_run = not same_column and not st.session_state.get("lqa_running", False)
    start_clicked = st.button(
        "开始校对",
        type="primary",
        disabled=not can_run,
        key="start_lqa",
    )

    result_df_existing = st.session_state.get("result_df")
    has_results = (
        result_df_existing is not None
        and not result_df_existing.empty
        and st.session_state.get("result_sheet") == selected_sheet
    )
    if has_results:
        orig_name = st.session_state.orig_name
        stem = orig_name.rsplit(".", 1)[0]
        out_name = f"{stem}_lqa_result.xlsx"
        buf = io.BytesIO()
        st.session_state.wb.save(buf)
        st.download_button(
            label="下载校对结果",
            data=buf.getvalue(),
            file_name=out_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            key="download_lqa_sidebar",
        )
    else:
        st.button("下载校对结果", disabled=True, key="download_lqa_placeholder")

    reset_clicked = st.button(
        "恢复原始文件",
        type="secondary",
        disabled=st.session_state.get("lqa_running", False),
        help="丢弃当前校对结果，回到刚上传时的工作簿。",
        key="reset_wb",
    )

# ---------------------------------------------------------------------------
# Actions
# ---------------------------------------------------------------------------
if reset_clicked and "orig_wb" in st.session_state:
    st.session_state.wb = clone_workbook(st.session_state.orig_wb)
    st.session_state.pop("result_df", None)
    st.session_state.pop("result_sheet", None)
    st.session_state.pop("lqa_error", None)
    st.session_state.pop("lqa_error_tb", None)
    st.rerun()

if start_clicked:
    st.session_state.lqa_cancel = False
    st.session_state.lqa_running = True
    st.session_state.pop("lqa_error", None)
    st.session_state.pop("lqa_error_tb", None)
    work_wb = clone_workbook(st.session_state.orig_wb)

    with st.spinner("正在进行 AI 校对，请稍候…（如需中止，请使用右上角菜单的 Stop）"):
        try:
            updated_wb = call_run_lqa(
                wb=work_wb,
                source_lang=source_lang,
                target_lang=target_lang,
                game_id=game_id,
                selected_sheet=selected_sheet,
                use_glossary=use_glossary,
            )

            st.session_state.wb = updated_wb
            updated_ws = updated_wb[selected_sheet]
            result_df = sheet_to_df(updated_ws)
            st.session_state.result_df = result_df
            st.session_state.result_sheet = selected_sheet
        except Exception as e:
            st.session_state.lqa_error = str(e)
            st.session_state.lqa_error_tb = traceback.format_exc()
            st.session_state.wb = clone_workbook(st.session_state.orig_wb)
        finally:
            st.session_state.lqa_running = False

    if "lqa_error" not in st.session_state:
        st.rerun()

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
result_df = st.session_state.get("result_df")
has_results = (
    result_df is not None
    and not result_df.empty
    and st.session_state.get("result_sheet") == selected_sheet
)

status = "校对完成" if has_results else "待校对"
render_hero(subtitle=f"{uploaded_file.name}  ·  {selected_sheet}  ·  {status}")

glossary_label = game_id if use_glossary else "未使用术语库"
render_meta(
    [
        uploaded_file.name,
        selected_sheet,
        f"{len(df):,} 行",
        f"{source_lang} → {target_lang}",
        glossary_label,
    ]
)

if st.session_state.get("lqa_error"):
    st.error(f"运行失败: {st.session_state.lqa_error}")
    with st.expander("查看详细错误"):
        st.code(st.session_state.lqa_error_tb or "")
elif has_results:
    st.success("校对完成，结果已写回工作簿。")

preview_rows = min(len(df), 200)
st.caption(f"共 {len(df)} 行，预览前 {preview_rows} 行。")
st.dataframe(df.head(preview_rows), use_container_width=True, height=520)
